"""
test_loader.py
--------------
This script creates four small sample files and tests core.loader.

Works from any cwd / IDE Run button.
"""

# bootstrap: 让脚本不依赖于"从 enterprise_rag/ 启动"
import sys
from pathlib import Path
sys.path.insert(0, str(Path(__file__).resolve().parent))

from docx import Document
from openpyxl import Workbook

from core.loader import load_documents


SAMPLE_DIR = Path("uploads") / "loader_test_samples"


def create_sample_files() -> list[Path]:
    """Create one sample file for each supported format."""

    SAMPLE_DIR.mkdir(parents=True, exist_ok=True)

    txt_path = SAMPLE_DIR / "sample.txt"
    txt_path.write_text("This is a TXT file.\nIt has two lines.", encoding="utf-8")

    md_path = SAMPLE_DIR / "sample.md"
    md_path.write_text(
        "# Markdown Title\n\nThis is **bold** text in a Markdown file.",
        encoding="utf-8",
    )

    docx_path = SAMPLE_DIR / "sample.docx"
    doc = Document()
    doc.add_paragraph("This is a Word paragraph.")
    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Name"
    table.cell(0, 1).text = "Score"
    table.cell(1, 0).text = "Alice"
    table.cell(1, 1).text = "98"
    doc.save(docx_path)

    xlsx_path = SAMPLE_DIR / "sample.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Employees"
    sheet.append(["Name", "Department"])
    sheet.append(["Alice", "Engineering"])
    sheet.append(["Bob", "Sales"])
    workbook.save(xlsx_path)

    return [docx_path, xlsx_path, txt_path, md_path]


def main() -> None:
    sample_files = create_sample_files()
    loaded_documents = load_documents(sample_files)

    for index, document in enumerate(loaded_documents, start=1):
        print("=" * 60)
        print(f"Document {index}")
        print("Metadata:")
        print(document["metadata"])
        print("\nContent preview:")
        print(document["content"][:300])


if __name__ == "__main__":
    main()
